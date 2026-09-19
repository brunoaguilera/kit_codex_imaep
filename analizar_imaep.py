#!/usr/bin/env python3
"""Extrae y analiza el nivel mensual del IMAEP total, serie original.

Primera entrega: limpieza, controles de calidad, estadisticas descriptivas,
grafica e informe. No realiza imputaciones, modelos ni pronosticos.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import statistics
from dataclasses import dataclass
from datetime import date, datetime
from importlib.metadata import version
from pathlib import Path
from typing import Iterable

os.environ.setdefault(
    "MPLCONFIGDIR", str(Path(__file__).resolve().parent / ".matplotlib-cache")
)

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from openpyxl import load_workbook


DEFAULT_INPUT = Path("data/raw/imaep_bcp_original.xlsx")
DEFAULT_PROCESSED_DIR = Path("data/processed")
DEFAULT_OUTPUTS_DIR = Path("outputs")
START = date(2016, 1, 1)
END = date(2025, 12, 1)
EXPECTED_OBSERVATIONS = 120
SOURCE_PAGE = (
    "https://www.bcp.gov.py/web/institucional/"
    "indicador-mensual-de-actividad-economica-del-paraguay-imaep-"
)
DOWNLOAD_URL = (
    "https://www.bcp.gov.py/documents/20117/2659103/"
    "Indicador%2BMensual%2Bde%2Bla%2BActividad%2BEcon%C3%B3mica%2Bdel%2BParaguay.xlsx/"
    "27688483-a192-9797-37f7-1f8b184bf430?t=1789069578102"
)


@dataclass(frozen=True)
class Observation:
    fecha: date
    imaep: float
    fecha_celda: str
    valor_celda: str
    preliminar: bool


@dataclass(frozen=True)
class WorkbookData:
    observations: list[Observation]
    base_year: int
    unit_label: str
    primary_range: str
    secondary_range: str
    secondary_equal: bool
    sheets: list[str]


class DataQualityError(RuntimeError):
    """Indica que los datos no permiten continuar el analisis con seguridad."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def month_sequence(start: date, end: date) -> list[date]:
    result: list[date] = []
    current = start
    while current <= end:
        result.append(current)
        year = current.year + (current.month // 12)
        month = current.month % 12 + 1
        current = date(year, month, 1)
    return result


def normalize_month(value: object, cell: str) -> date:
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    else:
        raise DataQualityError(f"{cell} no contiene una fecha valida: {value!r}")
    if result.day != 1:
        raise DataQualityError(f"{cell} no identifica el mes con dia 1: {result}")
    return result


def numeric_value(value: object, cell: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise DataQualityError(f"{cell} no contiene un numero: {value!r}")
    number = float(value)
    if not math.isfinite(number):
        raise DataQualityError(f"{cell} contiene un valor no finito: {value!r}")
    return number


def extract_base_year(label: object) -> int:
    if not isinstance(label, str):
        raise DataQualityError("IMAEP!B8 no contiene la unidad/base esperada.")
    normalized = " ".join(label.split())
    prefix = "Indice base "
    ascii_label = (
        normalized.replace("Í", "I").replace("í", "i").replace("É", "E").replace("é", "e")
    )
    if not ascii_label.startswith(prefix) or "= 100" not in ascii_label:
        raise DataQualityError(f"No se pudo interpretar la base en IMAEP!B8: {label!r}")
    year_text = ascii_label[len(prefix) :].split("=", 1)[0].strip()
    if not year_text.isdigit():
        raise DataQualityError(f"Ano base no interpretable en IMAEP!B8: {label!r}")
    return int(year_text)


def _extract_rows(ws, date_column: str, value_column: str) -> list[Observation]:
    observations: list[Observation] = []
    for row in range(1, ws.max_row + 1):
        date_cell = ws[f"{date_column}{row}"]
        if not isinstance(date_cell.value, (date, datetime)):
            continue
        month = normalize_month(date_cell.value, date_cell.coordinate)
        if START <= month <= END:
            value_cell = ws[f"{value_column}{row}"]
            observations.append(
                Observation(
                    fecha=month,
                    imaep=numeric_value(value_cell.value, value_cell.coordinate),
                    fecha_celda=date_cell.coordinate,
                    valor_celda=value_cell.coordinate,
                    preliminar=ws[f"A{row}"].value == "*",
                )
            )
    return observations


def validate_calendar(observations: Iterable[Observation]) -> tuple[list[str], list[str]]:
    dates = [item.fecha for item in observations]
    duplicates = sorted({item for item in dates if dates.count(item) > 1})
    expected = set(month_sequence(START, END))
    missing = sorted(expected - set(dates))
    return [item.isoformat() for item in duplicates], [item.isoformat() for item in missing]


def load_workbook_data(path: Path) -> WorkbookData:
    if not path.is_file():
        raise FileNotFoundError(f"No existe el Excel de entrada: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        required_sheets = {"IMAEP", "IMAEP-apertura por sectores"}
        missing_sheets = required_sheets - set(workbook.sheetnames)
        if missing_sheets:
            raise DataQualityError(f"Faltan hojas requeridas: {sorted(missing_sheets)}")

        primary = workbook["IMAEP"]
        if primary["C9"].value != "IMAEP Serie Original":
            raise DataQualityError(
                f"IMAEP!C9 no identifica la serie requerida: {primary['C9'].value!r}"
            )
        base_year = extract_base_year(primary["B8"].value)
        observations = _extract_rows(primary, "B", "C")
        duplicates, missing = validate_calendar(observations)
        if duplicates or missing or len(observations) != EXPECTED_OBSERVATIONS:
            raise DataQualityError(
                "Cobertura mensual invalida: "
                f"n={len(observations)}, duplicados={duplicates}, faltantes={missing}"
            )
        observations.sort(key=lambda item: item.fecha)

        secondary = workbook["IMAEP-apertura por sectores"]
        if secondary["O9"].value != "IMAEP" or secondary["O10"].value != "Serie Original":
            raise DataQualityError(
                "La verificacion cruzada no encontro IMAEP / Serie Original en O9:O10."
            )
        secondary_observations = _extract_rows(secondary, "B", "O")
        secondary_observations.sort(key=lambda item: item.fecha)
        secondary_pairs = [(item.fecha, item.imaep) for item in secondary_observations]
        primary_pairs = [(item.fecha, item.imaep) for item in observations]
        secondary_equal = primary_pairs == secondary_pairs
        if not secondary_equal:
            raise DataQualityError(
                "Los valores de IMAEP!C no coinciden con la serie original IMAEP de la hoja sectorial."
            )

        primary_range = (
            f"IMAEP!{observations[0].fecha_celda}:{observations[-1].valor_celda}"
        )
        secondary_range = (
            "IMAEP-apertura por sectores!"
            f"{secondary_observations[0].fecha_celda}:"
            f"{secondary_observations[-1].fecha_celda} y "
            f"{secondary_observations[0].valor_celda}:"
            f"{secondary_observations[-1].valor_celda}"
        )
        return WorkbookData(
            observations=observations,
            base_year=base_year,
            unit_label=f"Índice (base {base_year} = 100)",
            primary_range=primary_range,
            secondary_range=secondary_range,
            secondary_equal=secondary_equal,
            sheets=list(workbook.sheetnames),
        )
    finally:
        workbook.close()


def write_clean_csv(path: Path, observations: Iterable[Observation]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["fecha", "imaep"])
        for item in observations:
            writer.writerow([item.fecha.isoformat(), repr(item.imaep)])


def read_clean_csv(path: Path) -> list[tuple[date, float]]:
    result: list[tuple[date, float]] = []
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != ["fecha", "imaep"]:
            raise DataQualityError(f"Encabezado CSV inesperado: {reader.fieldnames}")
        for row in reader:
            month = date.fromisoformat(row["fecha"])
            value = float(row["imaep"])
            if month.day != 1 or not math.isfinite(value):
                raise DataQualityError(f"Fila CSV invalida: {row}")
            result.append((month, value))
    return result


def calculate_statistics(observations: list[Observation]) -> dict[str, object]:
    values = [item.imaep for item in observations]
    q1, _, q3 = statistics.quantiles(values, n=4, method="inclusive")
    minimum = min(values)
    maximum = max(values)
    variance = statistics.variance(values)
    standard_deviation = statistics.stdev(values)
    return {
        "cantidad_valida": len(values),
        "cantidad_faltante": EXPECTED_OBSERVATIONS - len(values),
        "media": statistics.fmean(values),
        "mediana": statistics.median(values),
        "minimo": minimum,
        "meses_minimo": [
            item.fecha.isoformat() for item in observations if item.imaep == minimum
        ],
        "maximo": maximum,
        "meses_maximo": [
            item.fecha.isoformat() for item in observations if item.imaep == maximum
        ],
        "rango": maximum - minimum,
        "q1": q1,
        "q3": q3,
        "rango_intercuartilico": q3 - q1,
        "varianza_muestral": variance,
        "desviacion_estandar_muestral": standard_deviation,
        "denominador_varianza": "n-1",
        "metodo_cuantiles": (
            "Cuantiles inclusivos con interpolación lineal; "
            "posición (n-1)*p (statistics.quantiles, method='inclusive')."
        ),
    }


def write_statistics_csv(path: Path, stats: dict[str, object]) -> None:
    rows = [
        ("cantidad_valida", str(stats["cantidad_valida"]), "observaciones", "", ""),
        ("cantidad_faltante", str(stats["cantidad_faltante"]), "meses", "", ""),
        ("media", f"{stats['media']:.3f}", "índice", "", ""),
        ("mediana", f"{stats['mediana']:.3f}", "índice", "", ""),
        (
            "minimo",
            f"{stats['minimo']:.3f}",
            "índice",
            ";".join(stats["meses_minimo"]),
            "",
        ),
        (
            "maximo",
            f"{stats['maximo']:.3f}",
            "índice",
            ";".join(stats["meses_maximo"]),
            "",
        ),
        ("rango", f"{stats['rango']:.3f}", "índice", "", ""),
        ("q1", f"{stats['q1']:.3f}", "índice", "", stats["metodo_cuantiles"]),
        ("q3", f"{stats['q3']:.3f}", "índice", "", stats["metodo_cuantiles"]),
        (
            "rango_intercuartilico",
            f"{stats['rango_intercuartilico']:.3f}",
            "índice",
            "",
            stats["metodo_cuantiles"],
        ),
        (
            "varianza_muestral",
            f"{stats['varianza_muestral']:.3f}",
            "índice^2",
            "",
            "Denominador n-1",
        ),
        (
            "desviacion_estandar_muestral",
            f"{stats['desviacion_estandar_muestral']:.3f}",
            "índice",
            "",
            "Denominador n-1",
        ),
    ]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["estadistica", "valor", "unidad", "meses_asociados", "metodo"])
        writer.writerows(rows)


def write_plot(path: Path, data: WorkbookData) -> None:
    dates = [item.fecha for item in data.observations]
    values = [item.imaep for item in data.observations]
    fig, ax = plt.subplots(figsize=(12, 6.5))
    ax.plot(dates, values, color="#1f4e79", linewidth=1.8)
    ax.set_title("Paraguay: IMAEP total, serie original mensual (2016-2025)")
    ax.set_xlabel("Mes")
    ax.set_ylabel(f"Índice (base {data.base_year} = 100)")
    ax.xaxis.set_major_locator(mdates.YearLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    ax.grid(axis="y", color="#d9d9d9", linewidth=0.7)
    ax.set_xlim(dates[0], dates[-1])
    fig.autofmt_xdate(rotation=0, ha="center")
    fig.text(
        0.01,
        0.01,
        "Fuente: Banco Central del Paraguay (BCP). Cifras preliminares, sujetas a revisión.",
        fontsize=8.5,
    )
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)


def write_report(
    path: Path,
    data: WorkbookData,
    stats: dict[str, object],
    source_hash: str,
    quality: dict[str, object],
) -> None:
    preliminary_count = sum(item.preliminar for item in data.observations)
    report = f"""# Análisis exploratorio del IMAEP mensual de Paraguay, 2016-2025

## Objetivo y fuente

Esta primera entrega describe el nivel mensual del Indicador Mensual de Actividad Económica del Paraguay (IMAEP) total, serie original, entre enero de 2016 y diciembre de 2025. El IMAEP ofrece una señal de corto plazo sobre la producción de bienes y servicios a precios constantes; no debe interpretarse como el valor del PIB mensual.

La fuente es el [Banco Central del Paraguay]({SOURCE_PAGE}). Se usó el libro oficial `imaep_bcp_original.xlsx` (SHA-256 `{source_hash}`). El archivo corresponde a una versión histórica consultada para este trabajo: sus cifras pueden incorporar revisiones respecto de publicaciones anteriores.

## Serie seleccionada y unidad

- Hoja: `IMAEP`.
- Encabezado: `C9`, **IMAEP Serie Original**.
- Fechas y valores: `{data.primary_range}`; fechas en la columna B y nivel en la columna C.
- Unidad: índice base {data.base_year} = 100, según `IMAEP!B8`.
- Frecuencia: mensual. En el CSV, el día `01` solo identifica cada mes y no implica frecuencia diaria.
- Verificación cruzada: `{data.secondary_range}`, encabezados `O9:O10` (**IMAEP / Serie Original**).

No se seleccionaron variaciones porcentuales, el IMAEP sin agricultura ni binacionales, sectores, serie ajustada ni tendencia-ciclo.

## Calidad de los datos

El período contiene **{stats['cantidad_valida']} observaciones válidas** y **{stats['cantidad_faltante']} meses faltantes**. Comienza en {quality['periodo_observado']['inicio']} y termina en {quality['periodo_observado']['fin']}; no se detectaron fechas duplicadas, valores no numéricos ni valores no finitos. Los {preliminary_count} registros del período tienen asterisco en el libro y el BCP los identifica como **cifras preliminares, sujetas a revisión**.

Todos los valores escritos en el CSV fueron releídos y comparados exactamente con las celdas del Excel. La segunda aparición de la serie en la hoja sectorial también coincide en los {stats['cantidad_valida']} meses. El detalle de cada correspondencia y de los controles está en [`calidad_datos.json`](calidad_datos.json).

## Serie temporal

![Serie temporal del IMAEP](serie_imaep.png)

La gráfica muestra fluctuaciones mensuales y niveles generalmente mayores hacia el final del período que al comienzo. El menor nivel observado fue **{stats['minimo']:.3f}** en **{', '.join(stats['meses_minimo'])}** y el mayor fue **{stats['maximo']:.3f}** en **{', '.join(stats['meses_maximo'])}**. Estas son descripciones del archivo; por sí solas no identifican causas ni prueban tendencia, estacionalidad o estacionariedad.

## Estadísticas básicas

| Estadística | Resultado |
|---|---:|
| Cantidad válida | {stats['cantidad_valida']} |
| Cantidad faltante | {stats['cantidad_faltante']} |
| Media | {stats['media']:.3f} |
| Mediana | {stats['mediana']:.3f} |
| Mínimo | {stats['minimo']:.3f} ({', '.join(stats['meses_minimo'])}) |
| Máximo | {stats['maximo']:.3f} ({', '.join(stats['meses_maximo'])}) |
| Rango | {stats['rango']:.3f} |
| Q1 | {stats['q1']:.3f} |
| Q3 | {stats['q3']:.3f} |
| Rango intercuartílico | {stats['rango_intercuartilico']:.3f} |
| Varianza muestral | {stats['varianza_muestral']:.3f} |
| Desviación estándar muestral | {stats['desviacion_estandar_muestral']:.3f} |

La varianza y la desviación estándar son muestrales y usan denominador **n-1**. Q1 y Q3 se calcularon mediante cuantiles inclusivos con interpolación lineal, posición `(n-1)*p`. Los cálculos usan los valores completos del Excel; el redondeo a tres decimales es solo de presentación. La tabla reutilizable está en [`estadisticas_basicas.csv`](estadisticas_basicas.csv).

## Conclusión y alcance

En el conjunto 2016-2025, la media fue **{stats['media']:.3f}**, la mediana **{stats['mediana']:.3f}** y el 50 % central de los niveles se ubicó entre **{stats['q1']:.3f}** y **{stats['q3']:.3f}**. El rango total fue **{stats['rango']:.3f}** puntos de índice. Estas medidas resumen niveles de una serie mensual preliminar; no constituyen un análisis causal ni un diagnóstico de estacionariedad.

Esta fase no incluye imputaciones, descomposición, ACF/PACF, pruebas de estacionariedad, modelos ni pronósticos.
"""
    path.write_text(report, encoding="utf-8")


def build_quality(
    input_path: Path,
    data: WorkbookData,
    csv_path: Path,
    csv_equal: bool,
    stats: dict[str, object],
) -> dict[str, object]:
    observations = data.observations
    duplicates, missing = validate_calendar(observations)
    correspondence = [
        {
            "fecha": item.fecha.isoformat(),
            "fila_csv": index + 2,
            "fecha_excel": f"IMAEP!{item.fecha_celda}",
            "valor_excel": f"IMAEP!{item.valor_celda}",
            "valor": item.imaep,
            "igual": True,
        }
        for index, item in enumerate(observations)
    ]
    variance_consistent = math.isclose(
        float(stats["desviacion_estandar_muestral"]) ** 2,
        float(stats["varianza_muestral"]),
        rel_tol=1e-12,
        abs_tol=1e-12,
    )
    return {
        "estado": "OK",
        "fuente": {
            "institucion": "Banco Central del Paraguay",
            "pagina": SOURCE_PAGE,
            "url_descarga": DOWNLOAD_URL,
            "archivo": str(input_path),
            "bytes": input_path.stat().st_size,
            "sha256": sha256_file(input_path),
            "hojas": data.sheets,
        },
        "seleccion": {
            "serie": "IMAEP total - Serie Original - nivel",
            "hoja": "IMAEP",
            "encabezado": "IMAEP!C9",
            "unidad_celda": "IMAEP!B8",
            "unidad": data.unit_label,
            "rango_primario": data.primary_range,
            "rango_verificacion": data.secondary_range,
        },
        "periodo_objetivo": {"inicio": START.isoformat(), "fin": END.isoformat()},
        "periodo_observado": {
            "inicio": observations[0].fecha.isoformat(),
            "fin": observations[-1].fecha.isoformat(),
        },
        "controles": {
            "observaciones_esperadas": EXPECTED_OBSERVATIONS,
            "observaciones_validas": len(observations),
            "fechas_duplicadas": duplicates,
            "meses_faltantes": missing,
            "valores_no_numericos": 0,
            "valores_no_finitos": 0,
            "orden_cronologico": all(
                observations[i].fecha < observations[i + 1].fecha
                for i in range(len(observations) - 1)
            ),
            "continuidad_mensual": not missing and not duplicates,
            "comparacion_hoja_sectorial": data.secondary_equal,
            "csv_releido": str(csv_path),
            "csv_igual_al_excel": csv_equal,
            "comparaciones_csv_excel": len(correspondence),
            "desviacion_estandar_cuadrada_aprox_varianza": variance_consistent,
            "cifras_preliminares": sum(item.preliminar for item in observations),
        },
        "correspondencia_csv_excel": correspondence,
        "versiones": {
            "python": __import__("platform").python_version(),
            "openpyxl": version("openpyxl"),
            "matplotlib": version("matplotlib"),
        },
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--processed-dir", type=Path, default=DEFAULT_PROCESSED_DIR)
    parser.add_argument("--outputs-dir", type=Path, default=DEFAULT_OUTPUTS_DIR)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    args.processed_dir.mkdir(parents=True, exist_ok=True)
    args.outputs_dir.mkdir(parents=True, exist_ok=True)
    quality_path = args.outputs_dir / "calidad_datos.json"
    try:
        data = load_workbook_data(args.input)
        csv_path = args.processed_dir / "imaep_2016_2025.csv"
        write_clean_csv(csv_path, data.observations)
        csv_rows = read_clean_csv(csv_path)
        excel_rows = [(item.fecha, item.imaep) for item in data.observations]
        csv_equal = csv_rows == excel_rows
        if not csv_equal:
            raise DataQualityError("El CSV releido no coincide exactamente con el Excel.")

        stats = calculate_statistics(data.observations)
        quality = build_quality(args.input, data, csv_path, csv_equal, stats)
        required_checks = quality["controles"]
        if not all(
            [
                required_checks["orden_cronologico"],
                required_checks["continuidad_mensual"],
                required_checks["comparacion_hoja_sectorial"],
                required_checks["csv_igual_al_excel"],
                required_checks["desviacion_estandar_cuadrada_aprox_varianza"],
            ]
        ):
            raise DataQualityError("Uno o mas controles obligatorios no pasaron.")

        write_statistics_csv(args.outputs_dir / "estadisticas_basicas.csv", stats)
        write_plot(args.outputs_dir / "serie_imaep.png", data)
        write_report(
            args.outputs_dir / "informe.md",
            data,
            stats,
            quality["fuente"]["sha256"],
            quality,
        )
        quality_path.write_text(
            json.dumps(quality, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    except Exception as exc:
        error_quality = {"estado": "ERROR", "error": str(exc), "tipo": type(exc).__name__}
        quality_path.write_text(
            json.dumps(error_quality, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(f"Analisis detenido: {exc}")
        return 1

    print(f"CSV limpio: {csv_path}")
    print(f"Calidad: {quality_path}")
    print(f"Grafica: {args.outputs_dir / 'serie_imaep.png'}")
    print(f"Estadisticas: {args.outputs_dir / 'estadisticas_basicas.csv'}")
    print(f"Informe: {args.outputs_dir / 'informe.md'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
